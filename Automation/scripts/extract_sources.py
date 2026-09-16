# -*- coding: utf-8 -*-
"""
extract_sources.py - สกัดข้อความจากไฟล์อ้างอิงทั้งหมดในโฟลเดอร์รายวิชา
ผลลัพธ์: _คลังข้อมูล/extracted/<relative path>.md  + manifest.json + INDEX.md

ใช้งาน (รันจากเครื่องไหนก็ได้ที่ sync โฟลเดอร์นี้จาก Google Drive):
    python "_คลังข้อมูล/tools/extract_sources.py"
    python "_คลังข้อมูล/tools/extract_sources.py" --force     # สกัดใหม่ทุกไฟล์

ต้องมี: python-docx, pymupdf (หรือ pypdf), pandas+openpyxl+xlrd, python-pptx
ไฟล์ PDF ที่สแกนมา (ไม่มีชั้นข้อความ) จะถูกทำเครื่องหมาย needs_ocr ใน manifest.json
"""
import argparse
import hashlib
import json
import os
import datetime

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
# คลังตำราอ้างอิงอยู่คนละโฟลเดอร์กับโฟลเดอร์งาน ระบุเพิ่มที่นี่ (path สัมพัทธ์กับ ROOT)
# ตัวอย่าง: os.path.abspath(os.path.join(ROOT, "..", "หนังสือ", "สอบสอน"))
# หรือส่งผ่านบรรทัดคำสั่ง --extra "<path>" (ใช้ซ้ำได้หลายครั้ง)
EXTRA_ROOTS = []
BASE = os.path.join(ROOT, "_คลังข้อมูล")
OUT = os.path.join(BASE, "extracted")
MANIFEST = os.path.join(BASE, "manifest.json")
INDEX = os.path.join(BASE, "INDEX.md")
SKIP_DIRS = {"_คลังข้อมูล", "__pycache__", ".git"}
EXTS = {".docx", ".pdf", ".xlsx", ".xls", ".txt", ".md", ".pptx"}


def retry_os(fn, tries=4, wait=3):
    """ไฟล์บน Google Drive อ่านครั้งแรกอาจ error เพราะยังสตรีมไม่เสร็จ ให้ลองซ้ำ"""
    import time
    for i in range(tries):
        try:
            return fn()
        except OSError:
            if i == tries - 1:
                raise
            time.sleep(wait)


def sha1(path, limit=8 * 1024 * 1024):
    h = hashlib.sha1()
    with open(path, "rb") as f:
        while True:
            b = f.read(1024 * 1024)
            if not b:
                break
            h.update(b)
            limit -= len(b)
            if limit <= 0:
                break
    return h.hexdigest()[:16]


def docx_to_md(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn
    d = docx.Document(path)
    out = []

    def blocks(parent):
        for ch in parent.element.body.iterchildren():
            if ch.tag == qn("w:p"):
                yield Paragraph(ch, parent)
            elif ch.tag == qn("w:tbl"):
                yield Table(ch, parent)

    for b in blocks(d):
        if isinstance(b, Paragraph):
            t = b.text.strip()
            if not t:
                continue
            st = (b.style.name or "").lower()
            if st.startswith("heading 1") or st == "title":
                out.append("## " + t)
            elif st.startswith("heading 2"):
                out.append("### " + t)
            elif st.startswith("heading"):
                out.append("#### " + t)
            elif st.startswith("list"):
                out.append("- " + t)
            else:
                out.append(t)
        else:
            rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in b.rows]
            if not rows:
                continue
            w = len(rows[0])
            out.append("")
            out.append("| " + " | ".join(rows[0]) + " |")
            out.append("|" + "---|" * w)
            for r in rows[1:]:
                out.append("| " + " | ".join(r) + " |")
            out.append("")
    return "\n\n".join(out), {}


def pdf_to_md(path):
    pages, empty = [], 0
    fitz = None
    try:
        import pymupdf as fitz
    except ImportError:
        try:
            import fitz
        except ImportError:
            fitz = None
    if fitz is not None:
        doc = fitz.open(path)
        n = doc.page_count
        for i, page in enumerate(doc, 1):
            t = page.get_text("text").strip()
            if len(t) < 30:
                empty += 1
                t = "*(หน้านี้ไม่มีชั้นข้อความ เป็นภาพสแกน ต้อง OCR หรือให้ผู้ช่วย AI อ่านจากภาพ)*"
            pages.append("### [หน้า %d]\n\n%s" % (i, t))
        doc.close()
    else:
        from pypdf import PdfReader
        r = PdfReader(path)
        n = len(r.pages)
        for i, p in enumerate(r.pages, 1):
            t = (p.extract_text() or "").strip()
            if len(t) < 30:
                empty += 1
                t = "*(หน้านี้ไม่มีชั้นข้อความ เป็นภาพสแกน)*"
            pages.append("### [หน้า %d]\n\n%s" % (i, t))
    meta = {"pages": n, "pages_no_text": empty, "needs_ocr": empty > 0 and empty >= n * 0.5}
    return "\n\n".join(pages), meta


def _df_to_md(df, limit=120):
    def cell(v):
        import math
        if v is None:
            return ""
        if isinstance(v, float) and math.isnan(v):
            return ""
        return str(v).replace("\n", " ").replace("|", "/").strip()
    rows = df.head(limit).values.tolist()
    if not rows:
        return "*(ไม่มีข้อมูล)*"
    w = max(len(r) for r in rows)
    out = ["| " + " | ".join("C%d" % (i + 1) for i in range(w)) + " |", "|" + "---|" * w]
    for r in rows:
        r = list(r) + [""] * (w - len(r))
        out.append("| " + " | ".join(cell(v) for v in r) + " |")
    if len(df) > limit:
        out.append("")
        out.append("*(แสดง %d จาก %d แถว)*" % (limit, len(df)))
    return "\n".join(out)


def xlsx_to_md(path):
    import pandas as pd
    out = []
    engine = "xlrd" if path.lower().endswith(".xls") else None
    sheets = None
    for eng in ([engine] if engine else [None]) + ["openpyxl", "xlrd", "calamine"]:
        try:
            sheets = pd.read_excel(path, sheet_name=None, header=None, engine=eng)
            break
        except Exception as e:
            last = e
    if sheets is None:
        # ไฟล์ .xls ที่จริงเป็น HTML table (ธปท. ส่งออกมาแบบนี้บ่อย)
        try:
            tables = pd.read_html(path)
            sheets = {"table_%d" % i: t for i, t in enumerate(tables, 1)}
        except Exception:
            return "*(อ่านไม่สำเร็จ: %s)*" % last, {"error": str(last)}
    for name, df in sheets.items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        out.append("### ชีต: %s  (%d แถว x %d คอลัมน์)" % (name, df.shape[0], df.shape[1]))
        out.append(_df_to_md(df))
    return "\n\n".join(out), {"sheets": list(sheets)}


def txt_to_md(path):
    for enc in ("utf-8", "cp874", "utf-16"):
        try:
            with open(path, encoding=enc) as f:
                return f.read(), {"encoding": enc}
        except (UnicodeDecodeError, UnicodeError):
            continue
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read(), {"encoding": "utf-8/replace"}


def pptx_to_md(path):
    from pptx import Presentation
    p = Presentation(path)
    out = []
    for i, s in enumerate(p.slides, 1):
        out.append("### สไลด์ %d" % i)
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text.strip())
    return "\n\n".join(out), {"slides": len(p.slides)}


HANDLERS = {".docx": docx_to_md, ".pdf": pdf_to_md, ".xlsx": xlsx_to_md,
            ".xls": xlsx_to_md, ".txt": txt_to_md, ".md": txt_to_md, ".pptx": pptx_to_md}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--extra", action="append", default=[], help="โฟลเดอร์คลังตำราเพิ่มเติม")
    a = ap.parse_args()
    EXTRA_ROOTS.extend(os.path.abspath(x) for x in a.extra)
    old = {}
    if os.path.exists(MANIFEST):
        with open(MANIFEST, encoding="utf-8") as f:
            old = {e["path"]: e for e in json.load(f)["files"]}
    entries = []
    walk_roots = [(ROOT, None)] + [(r, "คลังตำราอ้างอิง") for r in EXTRA_ROOTS if os.path.isdir(r)]
    for base_root, prefix in walk_roots:
      for dp, dns, fns in os.walk(base_root):
        dns[:] = [d for d in dns if d not in SKIP_DIRS and not d.startswith(".")]
        for fn in sorted(fns):
            ext = os.path.splitext(fn)[1].lower()
            if ext not in EXTS or fn.startswith("~$") or fn == "desktop.ini":
                continue
            full = os.path.join(dp, fn)
            rel = os.path.relpath(full, base_root).replace("\\", "/")
            if prefix:
                rel = prefix + "/" + rel
            st = os.stat(full)
            try:
                dig = retry_os(lambda: sha1(full))
            except OSError as e:
                # ไฟล์ Google Drive ที่ยังไม่ได้ดาวน์โหลดลงเครื่อง (เก็บบนคลาวด์อย่างเดียว)
                print("รอดาวน์โหลด", rel, e)
                entries.append({"path": rel, "chapter": rel.split("/")[0] if "/" in rel else "(root)",
                                "type": ext.lstrip("."), "bytes": st.st_size, "sha1": None,
                                "extracted_md": None, "chars": 0,
                                "error": "ยังไม่ได้ดาวน์โหลดจาก Google Drive (คลิกขวาไฟล์ > Offline access > Available offline)"})
                continue
            outp = os.path.join(OUT, rel + ".md")
            prev = old.get(rel)
            if prev and prev.get("sha1") == dig and os.path.exists(outp) and not a.force:
                entries.append(prev)
                print("skip", rel)
                continue
            os.makedirs(os.path.dirname(outp), exist_ok=True)
            err = None
            try:
                body, meta = retry_os(lambda: HANDLERS[ext](full))
            except OSError as e:
                print("รอดาวน์โหลด", rel, e)
                entries.append({"path": rel, "chapter": rel.split("/")[0] if "/" in rel else "(root)",
                                "type": ext.lstrip("."), "bytes": st.st_size, "sha1": dig,
                                "extracted_md": None, "chars": 0,
                                "error": "ยังไม่ได้ดาวน์โหลดจาก Google Drive"})
                continue
            except Exception as e:
                body, meta, err = "*(สกัดไม่สำเร็จ: %s)*" % e, {}, str(e)
            head = ("---\nsource_path: \"%s\"\nsha1: %s\nbytes: %d\nmodified: %s\nextracted: %s\n---\n\n# %s\n\n"
                    % (rel, dig, st.st_size,
                       datetime.datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"),
                       datetime.datetime.now().isoformat(timespec="seconds"), fn))
            with open(outp, "w", encoding="utf-8") as f:
                f.write(head + body)
            e = {"path": rel, "chapter": rel.split("/")[0] if "/" in rel else "(root)",
                 "type": ext.lstrip("."), "bytes": st.st_size, "sha1": dig,
                 "extracted_md": os.path.relpath(outp, BASE).replace("\\", "/"),
                 "chars": len(body)}
            e.update(meta)
            if err:
                e["error"] = err
            entries.append(e)
            print("ok  ", rel, e.get("chars"))
    os.makedirs(BASE, exist_ok=True)
    with open(MANIFEST, "w", encoding="utf-8") as f:
        json.dump({"generated": datetime.datetime.now().isoformat(timespec="seconds"),
                   "root": os.path.basename(ROOT), "count": len(entries), "files": entries},
                  f, ensure_ascii=False, indent=2)
    lines = ["# ดัชนีแหล่งข้อมูล", "",
             "สร้างอัตโนมัติโดย `_คลังข้อมูล/tools/extract_sources.py` เมื่อ %s"
             % datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), "",
             "ข้อความที่สกัดแล้วอยู่ในโฟลเดอร์ `extracted/` (โครงสร้างย่อยตรงกับไฟล์ต้นฉบับ)", ""]
    bych = {}
    for e in entries:
        bych.setdefault(e["chapter"], []).append(e)
    for ch in sorted(bych):
        lines += ["## %s" % ch, "",
                  "| ไฟล์ต้นฉบับ | ชนิด | ขนาด (KB) | ข้อความที่สกัด | หมายเหตุ |",
                  "|---|---|---|---|---|"]
        for e in sorted(bych[ch], key=lambda x: x["path"]):
            note = []
            if e.get("needs_ocr"):
                note.append("**สแกน ต้อง OCR**")
            elif e.get("pages_no_text"):
                note.append("บางหน้าไม่มีข้อความ (%s/%s)" % (e["pages_no_text"], e.get("pages", "?")))
            if e.get("pages"):
                note.append("%s หน้า" % e["pages"])
            if e.get("error"):
                note.append("ผิดพลาด: " + e["error"][:60])
            lines.append("| `%s` | %s | %d | `%s` | %s |"
                         % (os.path.basename(e["path"]), e["type"], e["bytes"] // 1024,
                            (e.get("extracted_md") or "-"), " / ".join(note)))
        lines.append("")
    with open(INDEX, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print("\nรวม %d ไฟล์ -> %s" % (len(entries), MANIFEST))


if __name__ == "__main__":
    main()
